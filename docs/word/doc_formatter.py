#!/usr/bin/env python3
"""
Módulo de formateo y renderizado para exportación de documentos Markdown a Word (.docx).
Universidad Popular del Cesar - Sistema PEA-i (Taller 2 Estructura de Datos).

Provee:
- Detección y renderizado de tablas con anchos proporcionales calculados, bordes sutiles,
  tblHeader (cabecera repetible en saltos de página), cantSplit (no partición de filas),
  sombreado alterno cebra y badges cromáticos de estado.
- Parser robusto de celdas Markdown con soporte de pipes dentro de código inline (`...|...`).
- Cajas de notas y citas (> ...) con barra lateral institucional izquierda y fondo tenue.
- Bloques de código preformateado (```) con marco y fondo de consola claro.
- Parseo de fórmulas matemáticas simples en línea ($ ... $).
- Tipografía institucional (Arial en encabezados, Calibri en cuerpo, Consolas en código).
"""

import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Colores institucionales UPC y temas complementarios
COLOR_PRIMARY_HEX = "1A365D"    # Azul institucional UPC (Header tablas, H1)
COLOR_SECONDARY_HEX = "0D9488"  # Verde UPC (H3, acentos)
COLOR_ACCENT_HEX = "2B6CB0"     # Azul medio (H2)
COLOR_ROW_ALT_HEX = "F8FAFC"    # Fondo alterno de fila (slate-50)
COLOR_ROW_WHITE_HEX = "FFFFFF"  # Blanco puro
COLOR_BORDER_HEX = "CBD5E1"     # Gris de bordes sutil (slate-300)
COLOR_BORDER_CALLOUT = "0D9488" # Verde UPC para borde izquierdo de callouts
COLOR_BG_CALLOUT = "F0FDFA"     # Fondo verde muy tenue para callouts
COLOR_BG_CODE = "F1F5F9"        # Fondo suave para bloques de código
COLOR_BORDER_CODE = "E2E8F0"    # Borde de bloque de código

# Configuración de Badges de estado con estilos visuales
BADGES_CONFIG = [
    (re.compile(r'cumplido', re.I), "DCFCE7", RGBColor(0x15, 0x80, 0x3D), "✓ Cumplido"),
    (re.compile(r'cubierto e implementado', re.I), "DCFCE7", RGBColor(0x15, 0x80, 0x3D), "✓ Cubierto (Impl)"),
    (re.compile(r'cubierto \(d-01\)', re.I), "DCFCE7", RGBColor(0x15, 0x80, 0x3D), "✓ Cubierto (D-01)"),
    (re.compile(r'cubierto \(implementado\)', re.I), "DCFCE7", RGBColor(0x15, 0x80, 0x3D), "✓ Cubierto (Impl)"),
    (re.compile(r'cubierto', re.I), "DCFCE7", RGBColor(0x15, 0x80, 0x3D), "✓ Cubierto"),
    (re.compile(r'pendiente', re.I), "FEF3C7", RGBColor(0xB4, 0x53, 0x09), "⏳ Pendiente"),
    (re.compile(r'en progreso', re.I), "DBEAFE", RGBColor(0x1D, 0x4E, 0xD8), "⚡ En progreso"),
    (re.compile(r'no aplica|n/a', re.I), "F1F5F9", RGBColor(0x64, 0x74, 0x8B), "– N/A –"),
    (re.compile(r'rechazado|error', re.I), "FEE2E2", RGBColor(0x99, 0x1B, 0x1B), "✕ Rechazado"),
]

def set_cell_background(cell, hex_color: str):
    """Aplica sombreado de fondo a una celda."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
    """Aplica márgenes internos (padding) a una celda en dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, border_color=COLOR_BORDER_HEX):
    """Configura bordes estilizados modernos: sin bordes verticales, líneas sutiles horizontales."""
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
        <w:left w:val="none"/>
        <w:bottom w:val="single" w:sz="12" w:space="0" w:color="{COLOR_PRIMARY_HEX}"/>
        <w:right w:val="none"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def make_row_header(row):
    """Marca la fila para repetirse como cabecera si la tabla salta de página."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))

def prevent_row_split(row):
    """Evita que una fila se divida horizontalmente a través de un salto de página."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))

def clean_latex_math(text: str) -> str:
    """Convierte expresiones LaTeX matemáticas frecuentes a texto limpio y legible para Word."""
    t = text
    t = re.sub(r'\$1900\s*\\le\s*\\text\{año\}\s*\\le\s*2100\$', '1900 ≤ año ≤ 2100', t)
    t = re.sub(r'\$start\s*\\le\s*end\$', 'start ≤ end', t)
    t = re.sub(r'\$([^\$]+)\$', r'\1', t)
    return t

def apply_inline_formatting(paragraph, text: str, default_color=None, default_size=Pt(9), is_header=False):
    """Parsea texto con negritas, código inline y matemáticas simples dentro de un párrafo."""
    clean_text = clean_latex_math(text)
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', clean_text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            inner = token[2:-2]
            r = paragraph.add_run(inner)
            r.font.name = "Arial" if is_header else "Calibri"
            r.font.bold = True
            r.font.size = default_size
            if default_color:
                r.font.color.rgb = default_color
        elif token.startswith('`') and token.endswith('`'):
            inner = token[1:-1]
            r = paragraph.add_run(inner)
            r.font.name = "Consolas"
            r.font.size = Pt(8.5) if default_size >= Pt(9) else Pt(8)
            if is_header:
                r.font.color.rgb = RGBColor(0xE0, 0xF2, 0xFE)
            else:
                r.font.color.rgb = RGBColor(0x03, 0x69, 0xA1) # Azul petróleo técnico
        else:
            r = paragraph.add_run(token)
            r.font.name = "Arial" if is_header else "Calibri"
            r.font.size = default_size
            if is_header:
                r.font.bold = True
            if default_color:
                r.font.color.rgb = default_color

def apply_cell_content(cell, text: str, is_header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    """Aplica formato enriquecido a una celda de tabla, incluyendo detección de badges."""
    cell.paragraphs[0].text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05

    raw_clean = text.strip()
    clean_no_md = raw_clean.replace('*', '').strip()

    # Detección de Badges de estado en celdas de datos
    if not is_header:
        for pattern, bg_hex, fg_rgb, badge_lbl in BADGES_CONFIG:
            if pattern.search(clean_no_md):
                r = p.add_run(f" {badge_lbl} ")
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.bold = True
                r.font.color.rgb = fg_rgb
                set_cell_background(cell, bg_hex)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                return

    text_color = RGBColor(0xFF, 0xFF, 0xFF) if is_header else RGBColor(0x33, 0x41, 0x55)
    text_size = Pt(9.5) if is_header else Pt(8.5)
    apply_inline_formatting(p, raw_clean, default_color=text_color, default_size=text_size, is_header=is_header)

def split_markdown_table_row(line: str) -> list[str]:
    """
    Divide una fila Markdown (| a | b | c |) respetando caracteres pipe (|)
    que se encuentren dentro de bloques de código inline (`...|...`).
    """
    chars = []
    in_backtick = False
    for ch in line:
        if ch == '`':
            in_backtick = not in_backtick
            chars.append(ch)
        elif ch == '|' and in_backtick:
            chars.append('__PIPE_SAFE__')
        else:
            chars.append(ch)

    masked = ''.join(chars).strip()
    if masked.startswith('|'):
        masked = masked[1:]
    if masked.endswith('|'):
        masked = masked[:-1]

    cells = [c.replace('__PIPE_SAFE__', '|').strip() for c in masked.split('|')]
    return cells

def calculate_dynamic_col_widths(headers: list[str], data_rows: list[list[str]], total_width_inches=6.5) -> list[Inches]:
    """
    Calcula anchos óptimos según el número de columnas y heurística del documento PEA-i.
    """
    num_cols = len(headers)
    if num_cols == 0:
        return []

    h_lower = [h.lower() for h in headers]
    first_col = h_lower[0] if h_lower else ""

    # Tabla 1: Sección 3 (Requisito, Descripción Taller, C++, Python/React, Estado)
    if num_cols == 5 and "requisito" in first_col and len(h_lower) > 1 and "descripci" in h_lower[1]:
        return [Inches(0.85), Inches(1.4), Inches(1.7), Inches(1.6), Inches(0.95)]

    # Tabla 2: Sección 5.1 (Módulo, Variable, Tipo, Restricción / Formato, Propósito)
    if num_cols == 5 and ("módulo" in first_col or "modulo" in first_col or "mdulo" in first_col):
        return [Inches(1.1), Inches(1.15), Inches(0.85), Inches(1.7), Inches(1.7)]

    # Tabla 3: Sección 28 (ID, Requisito, Componente, Evidencia esperada, Estado del diseño)
    if num_cols == 5 and first_col == "id":
        return [Inches(0.75), Inches(1.4), Inches(1.4), Inches(1.8), Inches(1.15)]

    # Tabla 4: Sección 32 (Estructura, Crear/Incluir, Consultar, Modificar, Desactivar, Eliminar/Extraer, Persistencia)
    if num_cols == 7 and "estructura" in first_col:
        return [Inches(1.15), Inches(0.85), Inches(0.85), Inches(0.85), Inches(0.85), Inches(0.95), Inches(1.0)]

    # Tabla 5: Sección 40 (Requisito, Regla/Entidad, Endpoint / Comando Real, Prueba, Evidencia)
    if num_cols == 5 and "requisito" in first_col and len(h_lower) > 1 and ("regla" in h_lower[1] or "entidad" in h_lower[1]):
        return [Inches(1.2), Inches(1.0), Inches(2.0), Inches(0.9), Inches(1.4)]

    # Cálculo genérico proporcional
    col_max_lengths = [len(h) for h in headers]
    for row in data_rows:
        for idx, val in enumerate(row):
            if idx < num_cols:
                col_max_lengths[idx] = max(col_max_lengths[idx], min(len(val), 60))

    total_len = sum(col_max_lengths) or 1
    widths = []
    min_width = 0.65
    for length in col_max_lengths:
        prop = length / total_len
        w = max(prop * total_width_inches, min_width)
        widths.append(w)

    scale = total_width_inches / sum(widths)
    return [Inches(round(w * scale, 2)) for w in widths]

def render_table_from_markdown(doc, clean_rows: list[list[str]]):
    """Renderiza una tabla limpia con todo el enriquecimiento institucional."""
    if not clean_rows:
        return

    headers = clean_rows[0]
    data_rows = clean_rows[1:]
    num_cols = len(headers)
    num_rows = len(clean_rows)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    col_widths = calculate_dynamic_col_widths(headers, data_rows)

    # 1. Configurar Encabezado
    header_row = table.rows[0]
    make_row_header(header_row)
    prevent_row_split(header_row)

    for c_idx, head_text in enumerate(headers):
        cell = header_row.cells[c_idx]
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_margins(cell, top=130, bottom=130, left=110, right=110)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        # Centrar si es columna de ID o Estado
        is_centered = (c_idx == 0 and len(head_text) <= 5) or ("estado" in head_text.lower())
        align = WD_ALIGN_PARAGRAPH.CENTER if is_centered else WD_ALIGN_PARAGRAPH.LEFT
        apply_cell_content(cell, head_text, is_header=True, align=align)

    # 2. Configurar Filas de datos
    for r_idx, row_values in enumerate(data_rows, start=1):
        row = table.rows[r_idx]
        prevent_row_split(row)
        bg_color = COLOR_ROW_ALT_HEX if (r_idx % 2 == 1) else COLOR_ROW_WHITE_HEX
        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            val = row_values[c_idx] if c_idx < len(row_values) else ""
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            is_centered = (c_idx == 0 and len(headers[0]) <= 5) or ("estado" in headers[c_idx].lower())
            align = WD_ALIGN_PARAGRAPH.CENTER if is_centered else WD_ALIGN_PARAGRAPH.LEFT
            apply_cell_content(cell, val, is_header=False, align=align)

    # 3. Anchos explícitos en cada celda
    for row in table.rows:
        for c_idx, width in enumerate(col_widths):
            if c_idx < len(row.cells):
                row.cells[c_idx].width = width

    # Espacio posterior sutil
    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(0)
    p_sep.paragraph_format.space_after = Pt(6)

def add_callout_box(doc, text: str, title: str = "NOTA TÉCNICA INSTITUCIONAL"):
    """
    Inserta una caja de llamada/nota elegante con borde verde institucional a la izquierda
    y fondo suave sombreado.
    """
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, COLOR_BG_CALLOUT)
    set_cell_margins(cell, top=110, bottom=110, left=150, right=140)

    # Borde grueso solo a la izquierda
    tblPr = tbl._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{COLOR_BORDER_CALLOUT}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
        <w:insideH w:val="none"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    if title:
        r_title = p.add_run(f"{title}\n")
        r_title.font.name = "Arial"
        r_title.font.bold = True
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)

    apply_inline_formatting(p, text, default_color=RGBColor(0x13, 0x4E, 0x4A), default_size=Pt(9))

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_text: str):
    """
    Inserta un bloque de código o diagrama ASCII dentro de un contenedor enmarcado
    con fondo de consola claro y fuente de ancho fijo.
    """
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, COLOR_BG_CODE)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    tblPr = tbl._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_CODE}"/>
        <w:left w:val="single" w:sz="14" w:space="0" w:color="{COLOR_ACCENT_HEX}"/>
        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_CODE}"/>
        <w:right w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_CODE}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.0)
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
