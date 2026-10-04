#!/usr/bin/env python3
"""
Generador del Documento Oficial Word (.docx) para la entrega del Taller 2 de Estructura de Datos.
Universidad Popular del Cesar - Grupo: WP-DN-EO
Procesa la especificación técnica maestra e incrusta imágenes retina de los diagramas y tablas estilizadas.
"""

import os
import re
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
SPEC_MD_PATH = REPO_ROOT / "docs" / "PEA-i_Especificacion_Integral_Cumplimiento_Taller2.md"
DIAGRAMS_DIR = REPO_ROOT / "docs" / "diagrams"
WORD_OUTPUT_DIR = REPO_ROOT / "docs" / "word"

# Iniciales oficiales
INITIALS = "WP-DN-EO"
DOCX_NAME = f"Taller2_{INITIALS}.docx"

# Colores institucionales UPC
COLOR_PRIMARY_HEX = "1A365D"    # Azul institucional oscuro
COLOR_SECONDARY_HEX = "0D9488"  # Verde UPC
COLOR_ACCENT_HEX = "2B6CB0"     # Azul medio
COLOR_BG_LIGHT_HEX = "F8FAFC"   # Fondo tenue
COLOR_CODE_BG_HEX = "0F172A"    # Fondo código
COLOR_BORDER_HEX = "CBD5E1"     # Borde tablas

def set_cell_background(cell, hex_color):
    """Aplica color de fondo a una celda de tabla."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Aplica márgenes internos (padding) a una celda."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_cover_page(doc):
    """Genera la portada institucional de la Universidad Popular del Cesar."""
    # Espaciado superior
    for _ in range(2):
        doc.add_paragraph()

    # Institución
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_inst.add_run("UNIVERSIDAD POPULAR DEL CESAR\n")
    run_inst.font.name = "Arial"
    run_inst.font.size = Pt(18)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    run_fac = p_inst.add_run("FACULTAD DE INGENIERÍAS Y TECNOLÓGICAS\nDEPARTAMENTO DE INGENIERÍA DE SISTEMAS\n")
    run_fac.font.name = "Arial"
    run_fac.font.size = Pt(13)
    run_fac.font.bold = True
    run_fac.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)

    run_sub = p_inst.add_run("ASIGNATURA: ESTRUCTURA DE DATOS\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    for _ in range(3):
        doc.add_paragraph()

    # Título del Proyecto
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("PEA-i: PROGRAMA ESTADÍSTICO DE ANÁLISIS DE INVESTIGACIÓN\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    run_taller = p_title.add_run("DOCUMENTO MAESTRO DE ESPECIFICACIÓN TÉCNICA, ARQUITECTURA,\nMODELO DE DATOS Y CUMPLIMIENTO DEL TALLER 2\n")
    run_taller.font.name = "Arial"
    run_taller.font.size = Pt(12)
    run_taller.font.bold = True
    run_taller.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    for _ in range(4):
        doc.add_paragraph()

    # Integrantes y Docente
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    run_doc_lbl = p_meta.add_run("DOCENTE:\n")
    run_doc_lbl.font.size = Pt(10)
    run_doc_lbl.font.bold = True
    run_doc_lbl.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    
    run_doc = p_meta.add_run("ING. ADITH PÉREZ\n\n")
    run_doc.font.size = Pt(12)
    run_doc.font.bold = True
    run_doc.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)

    run_int_lbl = p_meta.add_run("GRUPO DE TRABAJO (INICIALES):\n")
    run_int_lbl.font.size = Pt(10)
    run_int_lbl.font.bold = True
    run_int_lbl.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    run_int = p_meta.add_run(f"{INITIALS}\n(Grupo Oficial Taller 2 — 2026)\n")
    run_int.font.size = Pt(13)
    run_int.font.bold = True
    run_int.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)

    for _ in range(3):
        doc.add_paragraph()

    # Ciudad y Fecha
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_foot = p_foot.add_run("VALLEDUPAR, CESAR, COLOMBIA\nOCTUBRE DE 2026")
    run_foot.font.name = "Arial"
    run_foot.font.size = Pt(11)
    run_foot.font.bold = True
    run_foot.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    doc.add_page_break()

def setup_page_formatting(doc):
    """Configura márgenes estándar de 1 pulgada (2.54 cm)."""
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

        # Encabezado
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Universidad Popular del Cesar | PEA-i — Taller 2 (Estructura de Datos)")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

        # Pie de página
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun1 = fp.add_run(f"Grupo {INITIALS} — Especificación Técnica Integral")
        frun1.font.name = "Calibri"
        frun1.font.size = Pt(8.5)
        frun1.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

def insert_diagram(doc, diagram_key, caption_text, figure_num):
    """Inserta una imagen de diagrama de alta resolución centrada y con caption formal."""
    png_path = DIAGRAMS_DIR / f"{diagram_key}.png"
    if not png_path.exists():
        print(f"Advertencia: No se encontró {png_path}")
        return

    # Párrafo para la imagen
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    
    # Ajuste de ancho responsivo
    if "casos_de_uso" in diagram_key or "modelo_entidad" in diagram_key or "multilista" in diagram_key or "arquitectura" in diagram_key:
        run_img.add_picture(str(png_path), width=Inches(6.2))
    elif "hipercubo" in diagram_key or "clases_cpp" in diagram_key:
        run_img.add_picture(str(png_path), width=Inches(5.8))
    else:
        run_img.add_picture(str(png_path), width=Inches(5.5))

    # Párrafo del Pie de Figura (Caption)
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(14)
    
    run_lbl = p_cap.add_run(f"Figura {figure_num}: ")
    run_lbl.font.bold = True
    run_lbl.font.size = Pt(9.5)
    run_lbl.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    run_text = p_cap.add_run(caption_text)
    run_text.font.italic = True
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

def format_table(table, col_widths=None):
    """Estiliza una tabla Word con formato institucional UPC."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Encabezado (Fila 0)
    for cell in table.rows[0].cells:
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Filas de datos
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Anchos de columna si se proporcionan
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width

def parse_markdown_and_build_doc():
    """Parsea el documento maestro de especificación y ensambla el .docx completo."""
    print("Iniciando lectura de especificación...")
    content = SPEC_MD_PATH.read_text(encoding="utf-8")

    doc = docx.Document()
    setup_page_formatting(doc)
    create_cover_page(doc)

    lines = content.splitlines()
    in_code_block = False
    code_lang = ""
    code_lines = []
    
    in_table = False
    table_rows = []

    figure_counter = 1

    # Mapeo de secciones a imágenes de diagramas
    diagram_mapping = [
        (re.compile(r'### 38\.1 Diagrama de casos de uso', re.IGNORECASE), "01_casos_de_uso", "Diagrama de Casos de Uso del Sistema PEA-i (5 Actores y 6 Subsistemas)"),
        (re.compile(r'### 38\.2 Diagrama entidad-relación', re.IGNORECASE), "02_modelo_entidad_relacion", "Modelo Entidad-Relación Físico (16 Tablas SQL Server con Catálogo 2024)"),
        (re.compile(r'### 38\.3 Diagrama de clases C\+\+', re.IGNORECASE), "04_clases_cpp", "Diagrama de Clases C++ y Servicios del Núcleo de Dominio"),
        (re.compile(r'### 38\.3\.1 Esquema gráfico de punteros', re.IGNORECASE), "03_multilista_memoria_ram", "Disposición en Memoria RAM de la Multilista Ortogonal y Punteros C++"),
        (re.compile(r'#### Secuencia 1: Creación de Producto', re.IGNORECASE), "06_secuencia_creacion_escritura_dual", "Secuencia: Creación Write-Through en RAM y SQL Server con Pila LIFO"),
        (re.compile(r'#### Secuencia 2: Ingesta desde URL Scienti', re.IGNORECASE), "07_secuencia_ingesta_scienti", "Secuencia: Ingesta desde URL Scienti GrupLAC con Deduplicación Canónica"),
        (re.compile(r'#### Secuencia 3: Vinculación de Integrante', re.IGNORECASE), "08_secuencia_membresias_reglas", "Secuencia: Vinculación de Integrante y Validación de Reglas RV-001 y RV-002"),
        (re.compile(r'#### Secuencia 4: Validación Técnica de Evidencias', re.IGNORECASE), "09_secuencia_validacion_fifo", "Secuencia: Validación Técnica de Productos en Cola FIFO y Auditoría"),
        (re.compile(r'#### Secuencia 5: Deshacer Operación en Pila LIFO', re.IGNORECASE), "10_secuencia_deshacer_lifo", "Secuencia: Deshacer Última Operación en Pila LIFO (Reversión Simétrica)"),
        (re.compile(r'### 38\.5 Diagrama de despliegue', re.IGNORECASE), "11_arquitectura_despliegue", "Diagrama de Arquitectura de Integración y Despliegue Físico de Procesos"),
        (re.compile(r'### 38\.6 Diagrama del Hipercubo', re.IGNORECASE), "05_hipercubo_informacion", "Hipercubo de Información Multidimensional (4D/12D) y Operaciones OLAP")
    ]

    current_diagram_to_insert = None

    i = 0
    while i < len(lines):
        line = lines[i]

        # Detección de títulos de diagramas
        for pattern, diag_key, cap_text in diagram_mapping:
            if pattern.search(line):
                current_diagram_to_insert = (diag_key, cap_text)
                break

        # Manejo de bloques de código (```)
        if line.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = line[3:].strip().lower()
                code_lines = []
            else:
                in_code_block = False
                # Si era un bloque mermaid y teníamos un diagrama mapeado, insertamos la imagen renderizada
                if code_lang == "mermaid" and current_diagram_to_insert:
                    diag_key, cap_text = current_diagram_to_insert
                    insert_diagram(doc, diag_key, cap_text, figure_counter)
                    figure_counter += 1
                    current_diagram_to_insert = None
                else:
                    # Insertar bloque de código formateado
                    code_text = "\n".join(code_lines)
                    if code_text.strip():
                        p_code = doc.add_paragraph()
                        p_code.paragraph_format.space_before = Pt(6)
                        p_code.paragraph_format.space_after = Pt(8)
                        p_code.paragraph_format.left_indent = Inches(0.2)
                        run_code = p_code.add_run(code_text)
                        run_code.font.name = "Consolas"
                        run_code.font.size = Pt(8.5)
                        run_code.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Manejo de tablas Markdown (| Col 1 | Col 2 |)
        if line.strip().startswith("|") and line.strip().endswith("|"):
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(line)
            i += 1
            continue
        elif in_table:
            # Fin de la tabla: renderizarla
            in_table = False
            # Filtrar separadores (|---|---|)
            clean_rows = []
            for row_str in table_rows:
                parts = [c.strip() for c in row_str.strip().split("|")[1:-1]]
                if all(re.match(r'^:?-+:?$', p) for p in parts if p):
                    continue
                clean_rows.append(parts)

            if clean_rows:
                num_cols = max(len(r) for r in clean_rows)
                num_rows = len(clean_rows)
                table = doc.add_table(rows=num_rows, cols=num_cols)
                for r_idx, row_data in enumerate(clean_rows):
                    for c_idx, cell_value in enumerate(row_data):
                        if c_idx < num_cols:
                            table.cell(r_idx, c_idx).text = cell_value
                format_table(table)
                p_sep = doc.add_paragraph()
                p_sep.paragraph_format.space_after = Pt(8)

        # Encabezados Markdown
        if line.startswith("# "):
            heading_text = line[2:].strip()
            h = doc.add_heading(heading_text, level=1)
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(6)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.size = Pt(16)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            i += 1
            continue
        elif line.startswith("## "):
            heading_text = line[3:].strip()
            h = doc.add_heading(heading_text, level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.size = Pt(13.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            i += 1
            continue
        elif line.startswith("### "):
            heading_text = line[4:].strip()
            h = doc.add_heading(heading_text, level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.size = Pt(11.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
            i += 1
            continue
        elif line.startswith("#### "):
            heading_text = line[5:].strip()
            h = doc.add_heading(heading_text, level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.size = Pt(10.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            i += 1
            continue

        # Citas y notas (> ...)
        if line.startswith("> "):
            quote_text = line[2:].strip()
            p_q = doc.add_paragraph()
            p_q.paragraph_format.left_indent = Inches(0.3)
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(6)
            run_q = p_q.add_run(quote_text)
            run_q.font.italic = True
            run_q.font.size = Pt(9.5)
            run_q.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
            i += 1
            continue

        # Listas con viñetas (- o *)
        if re.match(r'^\s*[-*]\s+', line):
            indent_level = (len(line) - len(line.lstrip())) // 2
            bullet_text = re.sub(r'^\s*[-*]\s+', '', line).strip()
            # Eliminar checkboxes ([x] o [ ])
            bullet_text = re.sub(r'^\[[ xX]\]\s*', '', bullet_text)
            p_b = doc.add_paragraph(style='List Bullet')
            p_b.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            p_b.paragraph_format.space_before = Pt(1)
            p_b.paragraph_format.space_after = Pt(2)
            
            # Formatear negritas inline (**texto**)
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p_b.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p_b.add_run(part)
            i += 1
            continue

        # Listas numeradas (1. 2. etc.)
        if re.match(r'^\s*\d+\.\s+', line):
            num_text = re.sub(r'^\s*\d+\.\s+', '', line).strip()
            p_n = doc.add_paragraph(style='List Number')
            p_n.paragraph_format.space_before = Pt(1)
            p_n.paragraph_format.space_after = Pt(2)
            parts = re.split(r'(\*\*.*?\*\*)', num_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p_n.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p_n.add_run(part)
            i += 1
            continue

        # Líneas horizontales (---)
        if re.match(r'^-{3,}$', line.strip()):
            i += 1
            continue

        # Párrafos de texto regular
        stripped_line = line.strip()
        if stripped_line:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            # Parseo de negritas (**...**) y código inline (`...`)
            tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', stripped_line)
            for token in tokens:
                if token.startswith('**') and token.endswith('**'):
                    run = p.add_run(token[2:-2])
                    run.font.bold = True
                    run.font.name = "Calibri"
                    run.font.size = Pt(10)
                elif token.startswith('`') and token.endswith('`'):
                    run = p.add_run(token[1:-1])
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)
                else:
                    run = p.add_run(token)
                    run.font.name = "Calibri"
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        i += 1

    # Asegurar que la carpeta docs/word existe
    WORD_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_file = WORD_OUTPUT_DIR / DOCX_NAME
    doc.save(str(out_file))

    # También crear copia como WPDNEO.docx por conveniencia de entrega
    alt_file = WORD_OUTPUT_DIR / f"{INITIALS.replace('-', '')}.docx"
    doc.save(str(alt_file))

    size_mb = out_file.stat().st_size / (1024 * 1024)
    print(f"=== Documento Word generado con éxito ===")
    print(f"Archivo principal: {out_file} ({size_mb:.2f} MB)")
    print(f"Archivo alterno:   {alt_file}")

if __name__ == "__main__":
    parse_markdown_and_build_doc()
